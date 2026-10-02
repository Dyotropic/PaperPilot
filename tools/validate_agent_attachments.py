"""Business checks for local-file admission, restart/stop/compaction and vision.

Run with tools/run_validation.py; all fixtures and session assets stay isolated.
"""
import copy
import io
import json
from pathlib import Path
import unittest
from unittest.mock import patch

from PIL import Image
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from openpyxl import Workbook
import fitz

from paperpilot.agent_attachments import (AttachmentError, prepare_file, prepare_selection,
    persist_attachments, read_asset, format_attachment_material, anthropic_content,
    ensure_image_support, validate_request)
from paperpilot.conversation import ConversationManager
from paperpilot.context_budget import estimate_request_tokens
from paperpilot.agent_runtime import AgentRun, run_scope
from paperpilot.ai_service import AIService
from paperpilot.llm_client import LLMClient, ChatResult, AnthropicClient
from paperpilot.llm_usage import TokenUsage


class Client(LLMClient):
    def __init__(self):
        super().__init__("deepseek-flash")
        self.provider = "deepseek"
        self.requests = []

    def _do_chat(self, messages, *args):
        self.requests.append(copy.deepcopy(messages))
        compact = isinstance(messages[-1]["content"], str) and "现在生成科研工作流" in messages[-1]["content"]
        return ChatResult(content="## 用户目标与需求\n研究附件。\n## 关键数据与定位信息\n附件 result.csv 中参数为 42，图片为红色。" if compact else "已读取附件内容。",
                          finish_reason="stop", usage=TokenUsage(5000, 20, 1024, 3976))


class AttachmentBusiness(unittest.TestCase):
    def setUp(self):
        self.root = Path.cwd() / self._testMethodName
        self.root.mkdir()
        self.files = self.root / "source"
        self.files.mkdir()
        self.text = self.files / "result.csv"
        self.text.write_text("parameter,value\nx,42\n" + "synthetic evidence\n" * 160, encoding="utf-8")
        self.photo = self.files / "result.png"
        Image.new("RGB", (96, 80), "red").save(self.photo)
        self.cm = ConversationManager("fixture", session_id="abc", storage_path=self.root / "session" / "conversation.json")
        self.initialize(self.cm)

    def initialize(self, cm):
        cm.storage_directory.mkdir(parents=True, exist_ok=True)
        cm._path.with_name("events.jsonl").write_text(json.dumps(dict(kind="snapshot", data=cm._empty_data())) + "\n", encoding="utf-8")

    def refs(self):
        items, _ = prepare_selection([self.text, self.photo])
        return persist_attachments(self.cm.storage_directory, items)

    def test_snapshot_restart_and_compaction(self):
        refs = self.refs()
        self.cm.add_user_message(format_attachment_material(refs) + "读取参数和图片", display_content="读取参数和图片", attachments=refs)
        self.cm.add_assistant_message("x=42，图片为红色。")
        original = self.cm.build_api_messages("fixed system")
        self.text.unlink()
        self.photo.unlink()
        restored = ConversationManager("fixture", session_id="abc", storage_path=self.cm._path)
        self.assertEqual(restored.build_api_messages("fixed system"), original)
        self.assertIn("x,42", original[-2]["content"][0]["text"])
        self.assertTrue(original[-2]["content"][1]["image_url"]["url"].startswith("data:image/png;base64,"))
        self.assertNotIn("base64", restored._path.read_text(encoding="utf-8"))
        plan = restored.compaction_plan("fixed system", manual=True)
        self.assertTrue(restored.commit_compaction("附件 x=42，红色图片，继续研究。", plan,
                       mode="manual", provider="deepseek", model="deepseek-flash"))
        self.assertEqual(len(restored._history), 2)
        self.assertTrue(read_asset(restored.storage_directory, restored._history[0]["attachments"][0]))
        self.assertTrue(all(isinstance(m["content"], str) for m in restored.build_api_messages("fixed system")))

    def test_stop_and_abrupt_exit_before_user_event(self):
        refs = self.refs()
        run = AgentRun(self.cm, 0, "研究上传资料", attachments=refs)
        run.stop()
        run.finish()
        self.assertEqual(self.cm._messages[0]["attachments"], refs)
        self.assertIn("x,42", self.cm._messages[0]["content"])
        self.assertEqual(self.cm._meta["last_run"]["state"], "cancelled")
        other = ConversationManager("fixture", session_id="def", storage_path=self.root / "abrupt" / "conversation.json")
        self.initialize(other)
        refs = persist_attachments(other.storage_directory, prepare_selection([self.text, self.photo])[0])
        AgentRun(other, 0, "恢复图片研究", attachments=refs)
        restored = ConversationManager("fixture", session_id="def", storage_path=other._path)
        restored.recover_interrupted_run()
        self.assertEqual(restored._messages[0]["attachments"], refs)
        self.assertEqual(restored._meta["last_run"]["state"], "interrupted")

    def test_directory_scope_atomic_limits_and_excerpt(self):
        hidden = self.files / ".git"
        hidden.mkdir()
        (hidden / "secret.md").write_text("should not read", encoding="utf-8")
        (self.files / "binary.exe").write_bytes(b"ignored")
        items, note = prepare_selection(directory=self.files)
        self.assertEqual(len(items), 2)
        self.assertIn("跳过 2", note)
        self.assertTrue(all(i.name.startswith("source/") for i in items))
        for i in range(19):
            (self.files / f"more-{i}.txt").write_text("one file", encoding="utf-8")
        with self.assertRaisesRegex(AttachmentError, "超过 20"):
            prepare_selection(directory=self.files, existing=items)
        self.assertEqual(len(items), 2)
        self.text.write_text("A" * 21000, encoding="utf-8")
        excerpt = prepare_file(self.text)
        self.assertEqual(len(excerpt.excerpt), 20000)
        self.assertIn("20,000", excerpt.warnings[0])

    def test_documents_and_scanned_pages(self):
        doc = Document()
        run = doc.add_paragraph().add_run("Document parameter 42")
        run.font.name = "Microsoft YaHei"
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "Table evidence 51"
        doc.add_paragraph("After table evidence 52")
        for paragraph in doc.paragraphs:
            for run in paragraph.runs:
                run.font.name = "Microsoft YaHei"
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.font.name = "Microsoft YaHei"
        path = self.files / "article.docx"
        doc.save(path)
        document = prepare_file(path).excerpt
        self.assertIn("parameter 42", document)
        self.assertLess(document.index("Table evidence 51"), document.index("After table evidence 52"))
        slides = Presentation()
        slide = slides.slides.add_slide(slides.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text = "Slide evidence 73"
        path = self.files / "slides.pptx"
        slides.save(path)
        self.assertIn("evidence 73", prepare_file(path).excerpt)
        workbook = Workbook()
        workbook.active.append(["x", 42, "=SUM(B1:B1)"])
        path = self.files / "data.xlsx"
        workbook.save(path)
        item = prepare_file(path)
        self.assertIn("SUM", item.excerpt)
        self.assertIn("不执行", item.scope)
        pdf = fitz.open()
        pdf.new_page().insert_text((40, 60), "PDF evidence 87")
        path = self.files / "article.pdf"
        pdf.save(path)
        pdf.close()
        self.assertIn("evidence 87", prepare_file(path).excerpt)
        pdf = fitz.open()
        for _ in range(4):
            page = pdf.new_page(width=160, height=120)
            page.insert_image(page.rect, filename=str(self.photo))
        path = self.files / "scan.pdf"
        pdf.save(path)
        pdf.close()
        item = prepare_file(path)
        self.assertEqual(len(item.images), 3)
        self.assertIn("3/4", item.scope)

    def test_corrupt_missing_and_invalid_references(self):
        ref = self.refs()[1]
        asset = self.cm.storage_directory / "attachments" / ref["asset"]
        asset.write_bytes(b"corrupt")
        with self.assertRaisesRegex(AttachmentError, "完整性"):
            read_asset(self.cm.storage_directory, ref)
        asset.unlink()
        with self.assertRaisesRegex(AttachmentError, "缺失"):
            read_asset(self.cm.storage_directory, ref)
        with self.assertRaises(AttachmentError):
            read_asset(self.cm.storage_directory, dict(asset="../source/result.csv"))

    def test_request_limit_before_encoding_and_bounded_compaction(self):
        refs = self.refs()
        # Synthetic metadata models large historical image payloads without
        # allocating them. Admission must reject before reading/encoding bytes.
        large = copy.deepcopy(refs[1])
        large["images"][0]["size"] = 24 * 1024 * 1024
        self.cm.add_user_message("question", attachments=[large])
        self.cm.add_assistant_message("answer")
        with self.assertRaisesRegex(AttachmentError, "32 MiB"):
            self.cm.build_api_messages("system")
        self.cm.clear()
        for i in range(6):
            self.cm.add_user_message("question" * 20, attachments=[refs[1]])
            self.cm.add_assistant_message("answer")
        # Shrink the byte budget to exercise prefix selection on genuine small
        # images, while leaving room for the compression instruction reserve.
        from paperpilot.agent_attachments import estimated_request_bytes
        # A real oversized history can still compact a complete fitting prefix.
        for m in self.cm._messages:
            if m["role"] == "user":
                m["content"] += "x" * 5000
        with patch("paperpilot.agent_attachments.MAX_REQUEST_BYTES", 30_000):
            plan = self.cm.compaction_plan("system", keep_rounds=2)
        self.assertGreater(len(plan["batch"]), 0)
        self.assertLess(len(plan["batch"]), 8)
        self.assertEqual(plan["batch"][-1]["role"], "assistant")
        self.assertLess(estimated_request_bytes(plan["batch"]), 30_000)
        missing = self.root / "deleted-session"
        with self.assertRaisesRegex(AttachmentError, "目录已移除"):
            persist_attachments(missing, prepare_selection([self.text])[0])
        self.assertFalse(missing.exists())

    def test_real_chat_carriers_cache_prefix_and_unsupported_model(self):
        service = AIService()
        cm = service.get_conversation(0, "attachment business")
        refs = persist_attachments(cm.storage_directory, prepare_selection([self.text, self.photo])[0])
        client = Client()
        with patch("paperpilot.ai_service.get_client", return_value=client), \
             patch("paperpilot.ai_service.get_task_model", return_value="deepseek-flash"), \
             patch("paperpilot.ai_service.get_task_model_override", return_value=""):
            service.chat(0, "attachment business", "读取资料", session_id=cm.session_id, attachments=refs)
            first = copy.deepcopy(client.requests[-1])
            service.chat(0, "attachment business", "继续核对", session_id=cm.session_id)
            self.assertEqual(client.requests[-1][:len(first)], first)
            count = estimate_request_tokens(first)
            bigger = copy.deepcopy(first)
            bigger[-1]["content"][1]["image_url"]["url"] += "A" * 100000
            self.assertEqual(estimate_request_tokens(bigger), count)
            converted = anthropic_content(first[-1]["content"])
            self.assertEqual(converted[1]["type"], "image")
            self.assertEqual(converted[1]["source"]["media_type"], "image/png")
            self.assertEqual(AnthropicClient("test", "fixture")._split_messages(first)[1][-1]["content"], converted)
            self.assertEqual(service.compact_context(0, "attachment business", session_id=cm.session_id)["status"], "completed")
        with self.assertRaises(AttachmentError):
            ensure_image_support("deepseek", "deepseek-v4-pro")
        with self.assertRaises(AttachmentError):
            validate_request(first, "unknown", "unknown")


if __name__ == "__main__":
    unittest.main(verbosity=2)
